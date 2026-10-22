import datetime
from decimal import Decimal as D

from django.contrib.auth.models import User
from django.test import TestCase

from costing.models import Ingredient, Supplier
from money import calc
from money.models import Account, Allocation, Entry, SupplierPayment
from sales.models import CashUp, CashUpLine, Expense, ExpenseCategory, PaymentMethod
from stock.models import GRV, GRVLine

T = datetime.date(2026, 8, 3)


class MoneyTests(TestCase):
    def setUp(self):
        self.boss = User.objects.create_superuser("boss", "", "pw")
        self.client.force_login(self.boss)
        self.safe = Account.objects.create(name="Safe", kind=Account.SAFE, opening_date=T, opening_balance=D("400"))
        self.bank = Account.objects.create(name="Bank", kind=Account.BANK, opening_date=T, opening_balance=D("1000"))
        self.eco = Account.objects.create(name="EcoCash", kind=Account.MOBILE, opening_date=T, opening_balance=D("0"))
        self.sup = Supplier.objects.create(name="Aries", terms_days=30)
        self.flour = Ingredient.objects.create(name="Flour", unit="kg", pack_size=25, pack_price=D("20"), vat_rate=D("0.155"), supplier=self.sup)

    def grv(self, day, packs, price, invoice_total=None):
        g = GRV.objects.create(supplier=self.sup, date=day, invoice_no=f"INV{day.day}", invoice_total=invoice_total)
        GRVLine.objects.create(grv=g, ingredient=self.flour, packs=D(packs), pack_price=D(price))
        return g

    def test_cashup_flows_to_safe_bank_ecocash(self):
        cash = PaymentMethod.objects.get_or_create(name="Cash (US$)", defaults={"is_cash": True})[0]
        swipe = PaymentMethod.objects.get_or_create(name="Swipe / card")[0]
        eco = PaymentMethod.objects.get_or_create(name="EcoCash")[0]
        cu = CashUp.objects.create(date=T, opening_float=D("200"), cash_counted=D("700"), drops=D("500"))
        CashUpLine.objects.create(cashup=cu, method=cash, till_amount=D("1000"))
        CashUpLine.objects.create(cashup=cu, method=swipe, till_amount=D("300"), actual_amount=D("290"))
        CashUpLine.objects.create(cashup=cu, method=eco, till_amount=D("50"))
        bals = calc.balances(T)
        self.assertEqual(bals[self.safe], D("400") + D("700") + D("500") - D("200"))
        self.assertEqual(bals[self.bank], D("1290"))
        self.assertEqual(bals[self.eco], D("50"))

    def test_entries_and_transfer_and_expense(self):
        Entry.objects.create(date=T, account=self.bank, kind=Entry.OWNER_IN, amount=D("2000"))
        Entry.objects.create(date=T, account=self.safe, kind=Entry.TRANSFER, to_account=self.bank, amount=D("300"))
        cat = ExpenseCategory.objects.create(name="Gas")
        Expense.objects.create(date=T, category=cat, amount=D("50"), paid_from=Expense.BANK)
        Expense.objects.create(date=T, category=cat, amount=D("20"), paid_from=Expense.BANK, account=self.eco)
        Expense.objects.create(date=T, category=cat, amount=D("99"), paid_from=Expense.OWNER)
        bals = calc.balances(T)
        self.assertEqual(bals[self.bank], D("1000") + D("2000") + D("300") - D("50"))
        self.assertEqual(bals[self.safe], D("100"))
        self.assertEqual(bals[self.eco], D("-20"))
        cb = calc.cash_book(self.bank, T, T + datetime.timedelta(days=29))
        self.assertEqual(cb["closing"], D("3250"))
        bs = calc.balance_sheet(T)
        liab = dict(bs["liabilities"])
        self.assertEqual(liab["Owed to the owner / boss (paid from own pocket)"], D("99"))
        self.assertEqual(bs["capital"], D("2000"))

    def test_creditors_payment_and_ageing(self):
        g1 = self.grv(T, "10", "23.10")       # 231.00
        g2 = self.grv(T + datetime.timedelta(days=5), "5", "23.10", invoice_total=D("120"))
        later = T + datetime.timedelta(days=40)
        rows = calc.creditors(later)
        self.assertEqual([r["owing"] for r in rows], [D("231.00"), D("120")])
        self.assertEqual(rows[0]["late"], 10)
        r = self.client.post("/money/suppliers/pay/", {"supplier": self.sup.pk, "date": later.isoformat(), "account": self.bank.pk,
                                                      "amount": "300", f"grv_{g1.pk}": "231", f"grv_{g2.pk}": "69"})
        self.assertEqual(r.status_code, 302)
        p = SupplierPayment.objects.get()
        self.assertEqual(p.allocated, D("300"))
        rows = calc.creditors(later)
        self.assertEqual([(r["grv"].pk, r["owing"]) for r in rows], [(g2.pk, D("51"))])
        self.assertEqual(calc.balances(later)[self.bank], D("700"))
        st = calc.supplier_statement(self.sup, T, later)
        self.assertEqual(st["closing"], D("51"))

    def test_pay_must_add_up_and_not_overpay(self):
        g1 = self.grv(T, "1", "100")
        r = self.client.post("/money/suppliers/pay/", {"supplier": self.sup.pk, "date": T.isoformat(), "account": self.bank.pk,
                                                      "amount": "90", f"grv_{g1.pk}": "80"})
        self.assertContains(r, "must be the same")
        r = self.client.post("/money/suppliers/pay/", {"supplier": self.sup.pk, "date": T.isoformat(), "account": self.bank.pk,
                                                      "amount": "150", f"grv_{g1.pk}": "150"})
        self.assertContains(r, "at most")
        self.assertFalse(SupplierPayment.objects.exists())

    def test_owner_paid_supplier(self):
        g1 = self.grv(T, "1", "100")
        self.client.post("/money/suppliers/pay/", {"supplier": self.sup.pk, "date": T.isoformat(), "account": "owner",
                                                  "amount": "100", f"grv_{g1.pk}": "100"})
        self.assertEqual(calc.balances(T)[self.bank], D("1000"))
        bs = calc.balance_sheet(T)
        self.assertEqual(dict(bs["liabilities"])["Owed to the owner / boss (paid from own pocket)"], D("100"))

    def test_vat_worksheet(self):
        self.grv(T, "1", "115.50")            # VAT 15.50
        cat = ExpenseCategory.objects.create(name="Repairs")
        Expense.objects.create(date=T, category=cat, amount=D("23.10"), vat=D("3.10"), paid_from=Expense.BANK)
        Entry.objects.create(date=T, account=self.bank, kind=Entry.VAT, amount=D("5"))
        v = calc.vat_period(T, T)
        self.assertEqual(v["input"], D("18.60"))
        self.assertEqual(v["payable"], D("-18.60"))
        self.assertEqual(v["paid"], D("5"))

    def test_pages_open(self):
        self.grv(T, "1", "100")
        Entry.objects.create(date=T, account=self.bank, kind=Entry.LOAN_IN, amount=D("50"))
        for url in ["/money/", "/money/accounts/", f"/money/accounts/{self.bank.pk}/?month=2026-08", f"/money/accounts/{self.bank.pk}/check/",
                    "/money/entries/new/", "/money/suppliers/", f"/money/suppliers/{self.sup.pk}/", f"/money/suppliers/pay/?supplier={self.sup.pk}",
                    "/money/vat/?month=2026-08", "/money/balance-sheet/?date=2026-08-31"]:
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_bank_check(self):
        r = self.client.post(f"/money/accounts/{self.bank.pk}/check/", {"date": T.isoformat(), "statement_balance": "990", "notes": ""})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.bank.checks.get().difference, D("-10"))

    def test_rights(self):
        u = User.objects.create_user("cash", password="pw")
        self.client.force_login(u)
        self.assertEqual(self.client.get("/money/").status_code, 302)

    def test_no_accounts_yet(self):
        Account.objects.all().delete()
        self.assertContains(self.client.get("/money/"), "set up the money accounts")
        self.assertContains(self.client.get("/money/balance-sheet/"), "books start")


class PackTests(MoneyTests):
    def test_accountant_pack_has_money(self):
        self.grv(T, "1", "100")
        Entry.objects.create(date=T, account=self.bank, kind=Entry.OWNER_IN, amount=D("500"))
        r = self.client.get("/reports/accountant/?month=2026-08")
        self.assertContains(r, "Cash book – safe, bank")
        self.assertContains(r, "Balance sheet at")
        r = self.client.get("/reports/accountant/excel/?month=2026-08")
        self.assertEqual(r.status_code, 200)
        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        for name in ("Cash book summary", "Cash book", "Suppliers owed", "Balance sheet"):
            self.assertIn(name, wb.sheetnames)

    def test_grv_shows_paid_state(self):
        g = self.grv(T, "1", "100")
        self.assertContains(self.client.get(f"/stock/grv/{g.pk}/"), "Not paid yet")


class OpeningTests(TestCase):
    def test_what_was_there_at_the_start_is_not_profit(self):
        from equipment.models import EquipmentCategory as Category, Item
        T0 = datetime.date(2026, 11, 2)
        Account.objects.create(name="Bank", kind=Account.BANK, opening_date=T0, opening_balance=D("1000"))
        cat = Category.objects.create(name="Ovens", life_years=10)
        Item.objects.create(asset_no="EQ-1", description="Oven", category=cat, purchase_date=datetime.date(2026, 9, 1), unit_cost=D("12000"))
        bs = calc.balance_sheet(T0)
        self.assertEqual(bs["kept"], D("-100.00"))   # only November's wear and tear on the oven (12 000 over 10 years)
        self.assertEqual(bs["open_parts"]["money"], D("1000"))
