"""Excel downloads: every list opens as a real .xlsx, and the rights decide what is in it."""
import io
from decimal import Decimal as D

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase

from core.models import Access
from costing.models import Category, Ingredient, Recipe, RecipeLine, Supplier
from staff.models import Employee

from .export import EXPORTS


class ExcelTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_superuser("owner", password="TestPass123!")
        sup = Supplier.objects.create(name="Aries")
        flour = Ingredient.objects.create(name="Flour", supplier=sup, unit="kg", pack_size=D("25"), pack_price=D("20"))
        cat = Category.objects.create(name="Bread")
        r = Recipe.objects.create(number=1001, name="White loaf", category=cat, sell_price=D("1.20"), yield_override=10)
        RecipeLine.objects.create(recipe=r, ingredient=flour, qty=D("5"))
        Employee.objects.create(first_name="Ann", code="26001", rate=D("400"))

    def book(self, r):
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r["Content-Type"].startswith("application/vnd.openxmlformats"))
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def headers(self, ws):
        return [c.value for c in ws[4] if c.value]

    def test_owner_gets_every_list(self):
        self.client.force_login(self.owner)
        for key, item in EXPORTS.items():
            wb = self.book(self.client.get(f"/excel/{key}/"))
            self.assertTrue(wb.worksheets, key)
            if item[3]:
                self.book(self.client.get(f"/excel/{key}/?month=2026-10"))
        wb = self.book(self.client.get("/excel/products/"))
        ws = wb.active
        self.assertIn("GP %", self.headers(ws))
        self.assertEqual(ws["B5"].value, "White loaf")
        wb = self.book(self.client.get("/excel/recipes/"))
        self.assertEqual(wb.sheetnames, ["Recipes", "Ingredients per recipe"])
        self.assertEqual(wb["Ingredients per recipe"]["D5"].value, "Flour")
        wb = self.book(self.client.get("/excel/staff/"))
        self.assertEqual(wb["For the fingerprint machine"]["A5"].value, "26001")
        r = self.client.get("/excel/")
        self.assertContains(r, "Excel downloads")
        self.assertContains(r, "Ingredients")

    def test_rights(self):
        u = User.objects.create_user("till", password="x")
        Access.from_preset(u, "cashier")
        self.client.force_login(u)
        ws = self.book(self.client.get("/excel/products/")).active
        self.assertNotIn("GP %", self.headers(ws))
        self.assertNotIn("Cost per unit", self.headers(ws))
        for key in ("staff", "payrun", "ingredients", "recipes", "stock"):
            self.assertEqual(self.client.get(f"/excel/{key}/").status_code, 302, key)
        r = self.client.get("/excel/")
        self.assertNotContains(r, "/excel/staff/")
        self.assertContains(r, "/excel/products/")
        baker = User.objects.create_user("bake", password="x")
        acc = Access.from_preset(baker, "baker")
        acc.ingredients = 1  # may look at ingredients, but not at costs
        acc.save()
        self.client.force_login(baker)
        ws = self.book(self.client.get("/excel/ingredients/")).active
        self.assertNotIn("Pack price", self.headers(ws))
        self.assertEqual(self.client.get("/excel/products/").status_code, 302)
