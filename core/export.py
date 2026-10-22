"""Download lists as Excel files (.xlsx): ingredients, products & prices, recipes, stock, staff, sales …

Every download follows the rights: someone who may not see costs gets no cost columns, someone without
staff rights gets no staff list. The files have a brand-coloured header row, filters and totals that follow
the filter (SUBTOTAL), and the header stays at the top when scrolling.
"""
import datetime
import io
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone

from .access import rights
from .models import log_change

MONEY = '"$"#,##0.00'
MONEY4 = '"$"#,##0.0000'
NUM = '#,##0.00'
QTY = '#,##0.###'
INT = '#,##0'
PCT = '0.0%'
DATE = 'dd mmm yyyy'
WHEN = 'dd mmm yyyy hh:mm'

BURGUNDY, CREAM, PEACH, LINE, ORANGE = "7A1315", "FFF3D7", "FDEADC", "F4E4C9", "E56026"


class Col:
    def __init__(self, title, width=None, fmt=None, total=False):
        self.title, self.width, self.fmt, self.total = title, width, fmt, total


class Sheet:
    def __init__(self, name, cols, rows, note=""):
        self.name, self.cols, self.rows, self.note = name, cols, rows, note


def yes(v):
    return "yes" if v else "no"


def _month(request):
    try:
        return datetime.date.fromisoformat(request.GET.get("month", "") + "-01")
    except ValueError:
        return None


def _in_month(qs, field, m):
    if not m:
        return qs
    nxt = (m.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
    return qs.filter(**{f"{field}__gte": m, f"{field}__lt": nxt})


def _clean(v):
    if isinstance(v, datetime.datetime):
        v = timezone.localtime(v) if timezone.is_aware(v) else v
        return v.replace(tzinfo=None)
    if v is None or isinstance(v, (str, int, float, Decimal, datetime.date)):
        return v
    return str(v)


def workbook(title, sheets, who):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Pondo Bakery"
    wb.properties.title = title
    now = timezone.localtime()
    head_fill, head_font = PatternFill("solid", fgColor=BURGUNDY), Font(bold=True, color=CREAM)
    tot_fill, line = PatternFill("solid", fgColor=PEACH), Side(style="thin", color=LINE)
    for sh in sheets:
        ws = wb.create_sheet(sh.name[:31])
        n = len(sh.cols)
        ws["A1"] = f"Pondo Bakery – {sh.name}"
        ws["A1"].font = Font(bold=True, size=14, color=BURGUNDY)
        ws["A2"] = f"Exported {now:%d %b %Y %H:%M} by {who}" + (f" · {sh.note}" if sh.note else "") + " · amounts in US$"
        ws["A2"].font = Font(italic=True, size=9, color="7C5A4A")
        hr = 4
        for c, col in enumerate(sh.cols, 1):
            cell = ws.cell(row=hr, column=c, value=col.title)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[hr].height = 32
        widths = [max(len(col.title) * 0.9, 8) for col in sh.cols]
        r = hr
        for row in sh.rows:
            r += 1
            for c, (col, v) in enumerate(zip(sh.cols, row), 1):
                cell = ws.cell(row=r, column=c, value=_clean(v))
                if col.fmt:
                    cell.number_format = col.fmt
                cell.border = Border(bottom=line)
                if isinstance(v, str):
                    widths[c - 1] = max(widths[c - 1], min(len(v), 50) * 1.05)
                elif v is not None:
                    widths[c - 1] = max(widths[c - 1], 12)
        last = r
        if any(col.total for col in sh.cols) and last > hr:
            r += 1
            ws.cell(row=r, column=1, value="Total (follows the filter)").font = Font(bold=True, color=BURGUNDY)
            for c, col in enumerate(sh.cols, 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = tot_fill
                if col.total:
                    L = get_column_letter(c)
                    cell.value = f"=SUBTOTAL(109,{L}{hr + 1}:{L}{last})"
                    cell.number_format = col.fmt or NUM
                    cell.font = Font(bold=True)
        for c, col in enumerate(sh.cols, 1):
            ws.column_dimensions[get_column_letter(c)].width = col.width or min(max(widths[c - 1] + 2, 9), 55)
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)
        if last >= hr:
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(n)}{max(last, hr + 1)}"
        ws.page_setup.orientation = "landscape" if n > 6 else "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{hr}:{hr}"
        ws.sheet_properties.tabColor = ORANGE
        if not sh.rows:
            ws.cell(row=hr + 1, column=1, value="Nothing to show.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================================================================= the lists
def ex_ingredients(request, can):
    from costing.calc import Calculator
    from costing.models import Ingredient
    from django.db.models import Count

    calc = Calculator().preload()
    qs = (Ingredient.objects.select_related("supplier").annotate(n=Count("used_in__recipe", distinct=True))
          .order_by("-active", "name"))
    cols = [Col("Ingredient", 34), Col("Type"), Col("Supplier / group", 22), Col("Product code"), Col("Unit", 7),
            Col("Pack size", 10, QTY)]
    if can.costs:
        cols += [Col("Pack price", 12, MONEY4), Col("VAT in pack price", 10, PCT), Col("Cost per unit (excl. VAT)", 14, MONEY4)]
    cols += [Col("Re-order when below", 12, QTY), Col("Order up to", 12, QTY), Col("Used in recipes", 10, INT),
             Col("In use", 8), Col("Notes", 30)]
    rows = []
    for i in qs:
        row = [i.name, i.get_kind_display(), str(i.supplier or ""), i.code, i.unit, i.pack_size]
        if can.costs:
            row += [i.pack_price, i.vat_rate, calc.unit_cost(i)]
        row += [i.reorder_level, i.stock_par, i.n, yes(i.active), i.notes]
        rows.append(row)
    return "Ingredients", [Sheet("Ingredients", cols, rows, "archived ones at the bottom")]


def ex_products(request, can):
    from costing.calc import Calculator
    from costing.models import Recipe

    calc = Calculator().preload()
    qs = Recipe.objects.filter(active=True, is_subrecipe=False).select_related("category").order_by("category__sort_order", "category__name", "name")
    cols = [Col("No.", 8), Col("Product", 34), Col("Category", 18), Col("Sold per", 8),
            Col("Selling price incl. VAT", 13, MONEY), Col("VAT", 7, PCT), Col("Price excl. VAT", 13, MONEY)]
    if can.costs:
        cols += [Col("Cost per unit", 12, MONEY4), Col("GP $", 11, MONEY), Col("GP %", 8, PCT),
                 Col(f"Price for {calc.settings.target_gp:.0%} GP", 13, MONEY)]
    rows = []
    for r in qs:
        s = calc.summary(r)
        row = [r.number, r.name, str(r.category or ""), r.unit, r.sell_price,
               (s["vat"] if s["vat"] is not None else None), s["sell_excl"]]
        if can.costs:
            row += [s["cpu"], s["gp_amount"], s["gp"], s["target_price"]]
        rows.append(row)
    return "Products and prices", [Sheet("Products & prices", cols, rows)]


def ex_recipes(request, can):
    from costing.calc import Calculator, CycleError
    from costing.models import Recipe

    calc = Calculator().preload()
    recipes = list(Recipe.objects.select_related("category").order_by("-active", "number"))
    cols = [Col("No.", 8), Col("Recipe", 34), Col("Category", 18), Col("Mother recipe", 9), Col("Batch yield", 9, INT),
            Col("Dough weight (kg)", 11, QTY)]
    if can.costs:
        cols += [Col("Batch cost", 12, MONEY), Col("Cost per unit", 12, MONEY4)]
    if can.prices_view:
        cols += [Col("Selling price incl. VAT", 13, MONEY)]
    if can.costs:
        cols += [Col("GP %", 8, PCT)]
    cols += [Col("Baking temp", 12), Col("Baking time", 12), Col("In use", 8)]
    rows, lines = [], []
    for r in recipes:
        try:
            s = calc.summary(r)
        except CycleError:
            s = {}
        row = [r.number, r.name, str(r.category or ""), yes(r.is_subrecipe), s.get("yield"), s.get("dough")]
        if can.costs:
            row += [s.get("batch_cost"), s.get("cpu")]
        if can.prices_view:
            row += [r.sell_price]
        if can.costs:
            row += [s.get("gp")]
        row += [r.bake_temp, r.bake_time, yes(r.active)]
        rows.append(row)
        if not r.active:
            continue
        for ln in calc.lines(r):
            try:
                qty, cost = calc.line_qty(r, ln), (calc.line_cost(r, ln) if can.costs else None)
            except CycleError:
                qty, cost = None, None
            ln_row = [r.number, r.name, ln.get_section_display(), ln.ingredient.name, qty, ln.ingredient.unit]
            if can.costs:
                ln_row.append(cost)
            lines.append(ln_row)
    lcols = [Col("Recipe no.", 8), Col("Recipe", 30), Col("Part", 10), Col("Ingredient", 30), Col("Qty", 10, QTY), Col("Unit", 7)]
    if can.costs:
        lcols.append(Col("Cost", 11, MONEY4))
    return "Recipes", [Sheet("Recipes", cols, rows, "archived ones at the bottom"),
                       Sheet("Ingredients per recipe", lcols, lines, "one line per ingredient, for one batch")]


def ex_stock(request, can):
    from stock.levels import stock_rows

    rows_, lc = stock_rows()
    cols = [Col("Item", 34), Col("Supplier", 20), Col("On hand", 11, QTY), Col("Unit", 7), Col("Packs", 9, QTY),
            Col("Re-order when below", 12, QTY), Col("Order up to", 12, QTY), Col("Suggest order (packs)", 11, INT),
            Col("Low", 6)]
    if can.costs:
        cols.append(Col("Value", 12, MONEY, total=True))
    rows = []
    for x in rows_:
        i = x["i"]
        row = [i.name, str(i.supplier or ""), x["qty"], i.unit, x["packs"], i.reorder_level, i.stock_par, x["suggest"],
               yes(x["low"])]
        if can.costs:
            row.append(x["value"])
        rows.append(row)
    note = f"from the stock take of {lc.date:%d %b %Y} plus goods received, minus production" if lc else "no stock take yet"
    return "Stock levels", [Sheet("Stock levels", cols, rows, note)]


def ex_suppliers(request, can):
    from costing.models import Supplier

    cols = [Col("Supplier", 24), Col("Company name", 26), Col("Contact person", 20), Col("Cell / phone", 16),
            Col("Email", 24), Col("Address", 30), Col("Our account no.", 14), Col("Notes", 30), Col("In use", 8)]
    rows = [[s.name, s.company_name, s.contact, s.phone, s.email, s.address, s.account_no, s.notes, yes(s.active)]
            for s in Supplier.objects.order_by("-active", "name")]
    return "Suppliers", [Sheet("Suppliers", cols, rows)]


def ex_equipment(request, can):
    from equipment.models import Item

    today = timezone.localdate()
    cols = [Col("Asset no.", 10), Col("Item", 32), Col("Category", 18), Col("Location", 16), Col("Serial / model no.", 16),
            Col("Supplier", 16), Col("Bought on", 12, DATE), Col("Qty", 6, INT)]
    if can.costs:
        cols += [Col("Cost each (excl. VAT)", 12, MONEY), Col("Total cost", 12, MONEY, total=True),
                 Col("Useful life (years)", 9, NUM), Col("Depreciation a month", 12, MONEY, total=True),
                 Col("Book value now", 12, MONEY, total=True)]
    cols += [Col("Sold / scrapped on", 12, DATE), Col("What happened", 20)]
    rows = []
    for i in Item.objects.select_related("category", "location").order_by("asset_no"):
        row = [i.asset_no, i.description, str(i.category), str(i.location or ""), i.serial, i.supplier, i.purchase_date, i.qty]
        if can.costs:
            row += [i.unit_cost, i.total_cost, i.life, i.per_month, i.book_value(today)]
        row += [i.disposed_on, i.disposed_note]
        rows.append(row)
    return "Equipment", [Sheet("Equipment register", cols, rows)]


def ex_waste(request, can):
    from production.models import WasteLine

    m = _month(request)
    qs = _in_month(WasteLine.objects.select_related("log", "recipe", "log__recorded_by", "log__signed_off_by"), "log__date", m)
    cols = [Col("Date", 12, DATE), Col("Shift", 8), Col("Product", 30), Col("Qty", 8, QTY, total=True), Col("Reason", 14),
            Col("Notes", 24)]
    if can.costs:
        cols += [Col("Cost per unit", 12, MONEY4), Col("Cost", 12, MONEY, total=True)]
    cols += [Col("Written down by", 14), Col("Signed off by", 14)]
    rows = []
    for ln in qs.order_by("log__date", "log__shift", "pk"):
        row = [ln.log.date, ln.log.get_shift_display(), ln.recipe.name, ln.qty, ln.reason, ln.notes]
        if can.costs:
            row += [ln.unit_cost, ln.line_cost]
        row += [str(ln.log.recorded_by or ""), str(ln.log.signed_off_by or "")]
        rows.append(row)
    return "Waste", [Sheet("Waste", cols, rows, _month_note(m))]


def _month_note(m):
    return f"{m:%B %Y}" if m else "everything"


def ex_daily_sales(request, can):
    from costing.models import Category
    from sales.models import DailySales

    m = _month(request)
    days = list(_in_month(DailySales.objects.prefetch_related("lines"), "date", m).order_by("date"))
    cat_ids = sorted({ln.category_id for d in days for ln in d.lines.all()})
    cats = {c.pk: c.name for c in Category.objects.filter(pk__in=cat_ids)}
    order = sorted(cat_ids, key=lambda k: cats.get(k, ""))
    cols = [Col("Date", 12, DATE), Col("Day", 6)] + [Col(cats.get(k, "?"), 12, MONEY, total=True) for k in order]
    cols += [Col("Sales excl. VAT", 13, MONEY, total=True)]
    if can.costs:
        cols += [Col("GP %", 8, PCT), Col("Gross profit", 12, MONEY, total=True)]
    cols += [Col("Notes", 30)]
    rows = []
    for d in days:
        by = {ln.category_id: ln.amount for ln in d.lines.all()}
        row = [d.date, f"{d.date:%a}"] + [by.get(k) for k in order] + [d.total]
        if can.costs:
            row += [d.gp_pct, d.gp_amount]
        row.append(d.notes)
        rows.append(row)
    return "Daily sales", [Sheet("Daily sales", cols, rows, _month_note(m) + " · from the till's end-of-day report")]


def ex_product_sales(request, can):
    from sales.models import SaleLine

    m = _month(request)
    qs = _in_month(SaleLine.objects.select_related("day", "recipe"), "day__date", m).order_by("day__date", "position", "pk")
    cols = [Col("Date", 12, DATE), Col("No.", 8), Col("Product", 30), Col("Qty", 8, QTY, total=True),
            Col("Price each incl. VAT", 12, MONEY), Col("Value incl. VAT", 12, MONEY, total=True),
            Col("Value excl. VAT", 12, MONEY, total=True)]
    if can.costs:
        cols += [Col("Cost", 12, MONEY, total=True), Col("GP %", 8, PCT)]
    rows = []
    for ln in qs:
        row = [ln.day.date, ln.recipe.number, ln.recipe.name, ln.qty, ln.price, ln.value, ln.value_excl]
        if can.costs:
            row += [ln.cost, ln.gp]
        rows.append(row)
    return "Product sales", [Sheet("Product sales", cols, rows, _month_note(m))]


def ex_cashups(request, can):
    from sales.models import CashUp

    m = _month(request)
    qs = _in_month(CashUp.objects.select_related("signed_off_by"), "date", m).order_by("date", "till")
    cols = [Col("No.", 9), Col("Date", 12, DATE), Col("Till / shift", 12), Col("Cashier", 14),
            Col("Till total", 12, MONEY, total=True), Col("Over (+) / short (−)", 12, MONEY, total=True),
            Col("Float at the start", 11, MONEY), Col("Cash counted", 12, MONEY, total=True), Col("ZiG counted", 11, NUM),
            Col("ZiG rate", 9, NUM), Col("Float left", 11, MONEY), Col("Banked", 12, MONEY, total=True),
            Col("Signed off by", 14), Col("Signed off on", 16, WHEN), Col("Notes", 30)]
    rows = [[c.number, c.date, c.till, c.cashier, c.total_sales, c.total_variance, c.opening_float, c.cash_counted,
             c.zig_counted, c.zig_rate, c.float_left, c.banked, str(c.signed_off_by or ""), c.signed_off_at, c.notes]
            for c in qs]
    return "Cash-ups", [Sheet("Cash-ups", cols, rows, _month_note(m))]


def ex_expenses(request, can):
    from sales.models import Expense

    m = _month(request)
    qs = _in_month(Expense.objects.select_related("category", "cashup"), "date", m).order_by("date", "pk")
    cols = [Col("Date", 12, DATE), Col("Category", 18), Col("Paid to", 20), Col("What for", 30),
            Col("Amount", 12, MONEY, total=True), Col("VAT on the receipt", 11, MONEY, total=True),
            Col("Excl. VAT", 12, MONEY, total=True), Col("Paid from", 10), Col("Receipt / slip no.", 14), Col("From cash-up", 10)]
    rows = [[e.date, str(e.category), e.payee, e.description, e.amount, e.vat, e.amount_excl, e.get_paid_from_display(),
             e.receipt_no, e.cashup.number if e.cashup_id else ""] for e in qs]
    return "Expenses", [Sheet("Expenses", cols, rows, _month_note(m))]


def ex_invoices(request, can):
    from invoices.models import Document

    m = _month(request)
    qs = _in_month(Document.objects.all(), "date", m).order_by("date", "pk")
    cols = [Col("Number", 11), Col("Type", 16), Col("Date", 12, DATE), Col("Due / valid until", 12, DATE),
            Col("Customer", 22), Col("Company", 22), Col("Your ref / order no.", 14), Col("Status", 10),
            Col("Paid on", 12, DATE), Col("Paid how", 12), Col("Excl. VAT", 12, MONEY, total=True),
            Col("VAT", 11, MONEY, total=True), Col("Total incl. VAT", 13, MONEY, total=True)]
    rows = [[d.number, d.get_kind_display(), d.date, d.due, d.name, d.company, d.your_ref, d.get_status_display(),
             d.paid_on, d.paid_how, d.subtotal, d.vat, d.total] for d in qs]
    return "Invoices and quotes", [Sheet("Invoices & quotes", cols, rows, _month_note(m))]


def ex_deliveries(request, can):
    from invoices.models import Delivery

    m = _month(request)
    qs = _in_month(Delivery.objects.select_related("doc", "checked_by", "driver"), "date", m).order_by("date", "sent_at")
    cols = [Col("No.", 10), Col("Date", 12, DATE), Col("Left", 8), Col("Invoice", 11), Col("Customer", 22), Col("Cell", 14),
            Col("Address", 28), Col("Suburb", 16), Col("Driver no.", 10), Col("Driver", 18), Col("Driver cell", 14), Col("Car reg", 11),
            Col("Car", 16), Col("Goods incl. VAT", 12, MONEY, total=True), Col("Delivery fee (driver)", 12, MONEY, total=True),
            Col("Customer pays", 12, MONEY, total=True), Col("Paid how", 20), Col("Status", 12),
            Col("Money back from driver", 12, MONEY, total=True), Col("Short (−) / over (+)", 11, MONEY, total=True),
            Col("Rung up on till", 12, MONEY, total=True), Col("Till slip", 12), Col("Checked by", 12)]
    rows = []
    for d in qs:
        stage = {"out": "out", "delivered": "delivered", "done": "done", "cancelled": "cancelled"}[d.stage]
        rows.append([d.number, d.date, timezone.localtime(d.sent_at).strftime("%H:%M"), d.doc.number, d.doc.name,
                     d.doc.phone, d.doc.address1, d.doc.address2, d.driver.number if d.driver_id else "", d.driver_name, d.driver_phone, d.car_reg, d.car,
                     d.goods, d.fee, d.total, d.PAY_SHORT[d.pay], stage, d.money_back, d.difference, d.rung_up,
                     d.till_ref, str(d.checked_by or "")])
    return "Deliveries", [Sheet("Deliveries", cols, rows, _month_note(m))]


def ex_drivers(request, can):
    from invoices.drivers import _short_by_driver
    from invoices.models import Delivery, Driver
    from django.db.models import Count, Max, Q

    short = _short_by_driver()
    qs = Driver.objects.annotate(n=Count("deliveries", filter=~Q(deliveries__status=Delivery.CANCELLED)),
                                 last=Max("deliveries__date"))
    cols = [Col("Driver no.", 10), Col("Name", 22), Col("Cell", 14), Col("National ID no.", 16), Col("Driver's licence", 14),
            Col("Car reg", 11), Col("Make / model", 16), Col("Colour", 10), Col("ID card issued", 12, DATE),
            Col("Deliveries", 9, INT, total=True), Col("Last delivery", 12, DATE), Col("Short (−) / over (+)", 12, MONEY, total=True),
            Col("May deliver", 8), Col("Notes", 30)]
    rows = [[d.number, d.name, d.phone, d.id_number, d.licence_no, d.car_reg, d.car_make, d.car_colour, d.issued, d.n,
             d.last, short.get(d.pk), yes(d.active), d.notes] for d in qs]
    return "Drivers", [Sheet("Drivers", cols, rows)]


def ex_customers(request, can):
    from invoices.models import Customer

    cols = [Col("Name", 22), Col("Company", 24), Col("Address", 26), Col("Address line 2 / town", 20), Col("Phone", 16),
            Col("Email", 24), Col("VAT no.", 14), Col("Notes", 30), Col("In use", 8)]
    rows = [[c.name, c.company, c.address1, c.address2, c.phone, c.email, c.vat_no, c.notes, yes(c.active)]
            for c in Customer.objects.order_by("-active", "name")]
    return "Customers", [Sheet("Customers", cols, rows)]


def ex_staff(request, can):
    from staff.models import Employee, PayrollSettings, file_checklist

    needs = PayrollSettings.get().needs
    people = list(Employee.objects.prefetch_related("papers").order_by("-active", "code", "first_name"))
    cols = [Col("Staff no.", 10), Col("Known as", 12), Col("First name", 14), Col("Second name", 14), Col("Surname", 14), Col("National ID no.", 16),
            Col("Date of birth", 12, DATE), Col("Gender", 7), Col("Job", 22), Col("Department", 18), Col("Cell", 14),
            Col("Address", 28), Col("Next of kin", 24), Col("Started", 12, DATE), Col("Left on", 12, DATE),
            Col("Working here", 8), Col("Pay type", 13), Col("Salary / rate", 12, MONEY),
            Col("Transport allowance", 11, MONEY), Col("Other allowance", 11, MONEY), Col("Deduct NSSA", 8),
            Col("Deduct PAYE", 8), Col("NSSA no.", 14), Col("ZIMRA tax no.", 14), Col("Paid by", 12),
            Col("Bank / EcoCash details", 24), Col("Leave days a year", 9, NUM), Col("Staff file complete", 10)]
    rows = []
    for e in people:
        req = [i for i in file_checklist(e, needs, list(e.papers.all())) if i["required"]]
        done = sum(1 for i in req if i["state"] == "both" and i["exp"] != "expired")
        rows.append([e.code, e.known_as, e.first_name, e.middle_name, e.last_name, e.id_number, e.date_of_birth, e.gender, e.job_title,
                     e.get_department_display(), e.phone, e.address, e.next_of_kin, e.start_date, e.end_date, yes(e.active),
                     e.get_pay_type_display(), e.rate, e.transport_allowance, e.other_allowance, yes(e.pays_nssa),
                     yes(e.pays_paye), e.nssa_no, e.tax_no, e.get_paid_by_display(), e.bank_details, e.leave_days,
                     f"{done} of {len(req)}"])
    machine = [[e.code, e.name, e.get_department_display(), e.job_title] for e in people if e.active and e.code and not e.is_open]
    mcols = [Col("User ID (staff no.)", 14), Col("Name", 24), Col("Department", 20), Col("Job", 22)]
    return "Staff", [Sheet("Staff", cols, rows, "personal information – keep this file safe"),
                     Sheet("For the fingerprint machine", mcols, machine,
                           "use the staff no. as the user ID on the fingerprint machine")]


def ex_payrun(request, can):
    from staff.models import PayRun

    run = PayRun.objects.filter(pk=request.GET.get("run") or 0).first() or PayRun.objects.first()
    if not run:
        return "Pay run", [Sheet("Pay run", [Col("Nothing")], [])]
    cols = [Col("Staff no.", 10), Col("Name", 22), Col("Pay type", 12), Col("Rate", 11, MONEY), Col("Days worked", 8, NUM),
            Col("Hours worked", 8, NUM), Col("Unpaid days off", 8, NUM), Col("Overtime hours", 8, NUM),
            Col("Sunday / holiday hours", 9, NUM), Col("Basic", 12, MONEY, total=True),
            Col("Overtime pay", 11, MONEY, total=True), Col("Allowances", 11, MONEY, total=True),
            Col("Bonus", 11, MONEY, total=True), Col("Other earnings", 11, MONEY, total=True),
            Col("Gross", 12, MONEY, total=True), Col("NSSA", 11, MONEY, total=True), Col("PAYE", 11, MONEY, total=True),
            Col("AIDS levy", 10, MONEY, total=True), Col("Advances", 11, MONEY, total=True),
            Col("Other deductions", 11, MONEY, total=True), Col("Net pay", 12, MONEY, total=True),
            Col("NSSA – business", 11, MONEY, total=True), Col("ZIMDEF", 10, MONEY, total=True),
            Col("Paid by", 12), Col("Bank / EcoCash details", 24)]
    rows = []
    for ln in run.lines.select_related("employee"):
        e = ln.employee
        rows.append([e.code, e.name, ln.get_pay_type_display(), ln.rate, ln.days, ln.hours, ln.unpaid_days,
                     ln.overtime_hours, ln.holiday_hours, ln.basic, ln.overtime_pay, ln.allowances, ln.bonus,
                     ln.other_earnings, ln.gross, ln.nssa, ln.paye, ln.aids_levy, ln.advances, ln.other_deductions,
                     ln.net, ln.nssa_employer, ln.zimdef, e.get_paid_by_display(), e.bank_details])
    return f"Payroll {run.period:%Y-%m}", [Sheet(f"Payroll {run.period:%b %Y}", cols, rows,
                                                 f"{run.get_status_display()} · pay date {run.pay_date:%d %b %Y}")]


# key: (label, what is in it, group, needs a month?, who may download, builder)
EXPORTS = {
    "ingredients": ("Ingredients", "Every ingredient and packaging item with supplier, pack size and (if you may see costs) prices",
                    "Bakery & Stock", False, lambda c: c.ingredients_view, ex_ingredients),
    "recipes": ("Recipes", "All recipes, plus a second sheet with the ingredients of every recipe",
                "Bakery & Stock", False, lambda c: c.recipes_view, ex_recipes),
    "products": ("Products & prices", "Every product on sale with its selling price (and cost and GP if you may see costs)",
                 "Bakery & Stock", False, lambda c: c.prices_view, ex_products),
    "stock": ("Stock levels", "What is in stock now and what to order", "Bakery & Stock", False,
              lambda c: c.stock_view, ex_stock),
    "suppliers": ("Suppliers", "Names, contacts and account numbers", "Bakery & Stock", False,
                  lambda c: c.stock_view, ex_suppliers),
    "equipment": ("Equipment register", "Every item with asset number, location (and cost and book value)",
                  "Bakery & Stock", False, lambda c: c.equipment_view, ex_equipment),
    "waste": ("Waste", "Every product written off, with the reason", "Bakery & Stock", True,
              lambda c: c.waste_view, ex_waste),
    "daily_sales": ("Daily sales", "Sales per day and per category, from the till's end-of-day report", "Sales & money",
                    True, lambda c: c.sales_view, ex_daily_sales),
    "product_sales": ("Product sales", "How many of each product was sold per day", "Sales & money", True,
                      lambda c: c.sales_view, ex_product_sales),
    "cashups": ("Cash-ups", "Every till cash-up with over / short and banking", "Sales & money", True,
                lambda c: c.sales_view, ex_cashups),
    "expenses": ("Expenses", "Every expense and pay-out with VAT", "Sales & money", True,
                 lambda c: c.sales_view, ex_expenses),
    "invoices": ("Invoices & quotes", "Every invoice, pro-forma and quotation with totals and status", "Sales & money",
                 True, lambda c: c.invoices_view, ex_invoices),
    "deliveries": ("Deliveries", "Every delivery with driver, car, fee and money back", "Sales & money", True,
                   lambda c: c.deliveries_view, ex_deliveries),
    "drivers": ("Delivery drivers", "Every driver with ID, cell, car and how often the money was short", "Sales & money",
                False, lambda c: c.deliveries_view, ex_drivers),
    "customers": ("Customers", "Names, contacts and VAT numbers", "Sales & money", False,
                  lambda c: c.invoices_view, ex_customers),
    "staff": ("Staff", "All staff details and pay, plus a sheet for the fingerprint machine", "Staff", False,
              lambda c: c.staff_view, ex_staff),
    "payrun": ("Pay run", "Every line of one month's payroll", "Staff", False, lambda c: c.staff_view, ex_payrun),
}


@login_required
def export_file(request, key):
    item = EXPORTS.get(key)
    can = rights(request)
    if not item or not item[4](can):
        messages.error(request, "You don't have rights for that list. Ask the Manager if you need it.")
        return redirect("core:dashboard")
    name, sheets = item[5](request, can)
    data = workbook(name, sheets, request.user.get_full_name() or request.user.username)
    m = _month(request) if item[3] else None
    fname = f"Pondo {name}{' ' + format(m, '%Y-%m') if m else ''} {timezone.localdate():%Y-%m-%d}.xlsx"
    log_change(request.user, "Excel", item[0], "downloaded", _month_note(m) if item[3] else "")
    resp = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp


@login_required
def export_hub(request):
    from staff.models import PayRun

    can = rights(request)
    groups = {}
    for key, (label, about, group, monthly, ok, _) in EXPORTS.items():
        if key != "payrun" and ok(can):
            groups.setdefault(group, []).append({"key": key, "label": label, "about": about, "monthly": monthly})
    runs = list(PayRun.objects.all()[:12]) if can.staff_view else []
    return render(request, "core/export.html", {"groups": groups, "runs": runs,
                                                "this_month": f"{timezone.localdate():%Y-%m}"})
