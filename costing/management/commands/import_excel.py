"""Import suppliers, categories, ingredients, recipes and selling prices from PB_Master_File.xlsx.

    python manage.py import_excel "C:\\path\\PB_Master_File.xlsx" [--replace]

Reads the values Excel last calculated, plus the formulas (to keep "per item" quantities live).
"""
import re
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import CompanySettings, Shop
from costing.calc import Calculator
from costing.models import Category, Ingredient, Recipe, RecipeLine, Supplier

UNIT_MAP = {"kg": "kg", "litre": "L", "liter": "L", "l": "L", "each": "each", "sheet": "sheet", "metre": "m",
            "meter": "m", "m": "m"}
NOT_SUPPLIERS = {"mother recipe", "packaging"}


def D(v):
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v))
    except (InvalidOperation, ValueError):
        return None


def clean(s):
    return re.sub(r"\s+", " ", str(s)).strip() if s is not None else ""


class Command(BaseCommand):
    help = "Import recipes & costing data from the Pondo master Excel file"

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--replace", action="store_true", help="Delete existing recipes/ingredients first")

    def handle(self, xlsx, replace=False, **opts):
        import openpyxl

        self.stdout.write("Reading the Excel file (this takes a minute)…")
        wf = openpyxl.load_workbook(xlsx, read_only=True)  # formulas
        wv = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)  # values
        need = ["Categories", "Ingredient Cost", "Recipe Database", "Product Cost & Sell"]
        for n in need:
            if n not in wv.sheetnames:
                raise CommandError(f"Sheet '{n}' not found – is this the Pondo master file?")
        if Recipe.objects.exists() and not replace:
            raise CommandError("There are already recipes in the program. Use --replace to start again from Excel.")

        def rows(wb, sheet, first, last, maxcol):
            return [list(r) for r in wb[sheet].iter_rows(min_row=first, max_row=last, max_col=maxcol, values_only=True)]

        cat_v = rows(wv, "Categories", 8, 220, 5)
        ing_v = rows(wv, "Ingredient Cost", 7, 600, 11)
        ing_f = rows(wf, "Ingredient Cost", 7, 600, 11)
        rd_v = rows(wv, "Recipe Database", 6, 1002, 94)
        rd_f = rows(wf, "Recipe Database", 6, 1002, 94)
        ps_v = rows(wv, "Product Cost & Sell", 6, 1002, 11)
        sup_v = rows(wv, "Suppliers", 1, 80, 6) if "Suppliers" in wv.sheetnames else []
        shop_v = {clean(r[0]).lower(): r[1] for r in sup_v if r and r[0]}

        report = {"warnings": []}
        with transaction.atomic():
            if replace:
                RecipeLine.objects.all().delete()
                Ingredient.objects.all().delete()
                Recipe.objects.all().delete()
            # ---- shop details from the Suppliers sheet
            shop = Shop.objects.first() or Shop()
            shop.name = clean(shop_v.get("shop name")) or "Pondo Bakery"
            shop.address_line1 = clean(shop_v.get("address line 1"))
            shop.address_line2 = clean(shop_v.get("address line 2 / town"))
            shop.phone = clean(shop_v.get("phone"))
            shop.email = clean(shop_v.get("email"))
            shop.tin = clean(shop_v.get("tin. no."))
            shop.save()
            CompanySettings.get()

            # ---- categories and suppliers
            cats = {}
            for i, r in enumerate(cat_v):
                n = clean(r[0])
                if n:
                    cats[n.lower()], _ = Category.objects.get_or_create(name=n, defaults={"sort_order": i})
            sups = {}

            def supplier(name):
                n = clean(name)
                if not n or n.lower() in NOT_SUPPLIERS:
                    return None
                if n.lower() not in sups:
                    sups[n.lower()], _ = Supplier.objects.get_or_create(name=n)
                return sups[n.lower()]

            for r in cat_v:
                supplier(r[4])

            # ---- recipes (first pass: the recipe itself)
            recipes_by_id, recipes_by_row = {}, {}
            for idx, (rv, rf) in enumerate(zip(rd_v, rd_f)):
                rid, name = rv[0], clean(rv[1])
                if not name or not isinstance(rid, (int, float)):
                    continue
                if name.lower() in {r.name.lower() for r in recipes_by_id.values()}:
                    report["warnings"].append(f"Recipe name used twice, kept the first: {name}")
                    continue
                cat_name = clean(rv[2])
                cat = cats.get(cat_name.lower()) if cat_name else None
                if cat_name and not cat:
                    cat = cats[cat_name.lower()] = Category.objects.create(name=cat_name, sort_order=999)
                temp = rv[4]
                rec = Recipe.objects.create(
                    number=201000 + int(rid), name=name, category=cat,
                    is_subrecipe=(cat_name.lower() == "mother recipe"),
                    scaling_weight=D(rv[7]) if isinstance(rv[7], (int, float)) and rv[7] else None,
                    weight_loss=D(rv[93]) or Decimal("0"),
                    bake_temp=clean(int(temp) if isinstance(temp, float) and temp.is_integer() else temp),
                    bake_time=clean(rv[5]), method=(rv[8] or "").strip() if isinstance(rv[8], str) else "",
                    allergens=(rv[9] or "").strip() if isinstance(rv[9], str) else "",
                )
                # selling price from Product Cost & Sell (same row)
                pv = ps_v[idx] if idx < len(ps_v) else None
                if pv and isinstance(pv[6], (int, float)):
                    rec.sell_price = D(round(pv[6], 2))
                    rec.vat_rate = D(pv[7]) if isinstance(pv[7], (int, float)) else None
                    rec.save()
                recipes_by_id[int(rid)] = rec
                recipes_by_row[idx + 6] = rec

            # ---- ingredients
            names = {}
            for rv, rf in zip(ing_v, ing_f):
                name = clean(rv[0])
                if not name:
                    continue
                if name.lower() in names:
                    report["warnings"].append(f"Ingredient listed twice, kept the first: {name}")
                    continue
                group = clean(rv[1])
                unit = UNIT_MAP.get(clean(rv[6]).lower(), "kg")
                formula = rf[3] if isinstance(rf[3], str) and rf[3].startswith("=") else ""
                linked = None
                m = re.search(r"MATCH\((\d+),'Recipe Database'", formula)
                if m:
                    linked = recipes_by_id.get(int(m.group(1)))
                m2 = re.search(r"'Product Cost & Sell'!\$?F\$?(\d+)", formula)
                if m2:
                    linked = recipes_by_row.get(int(m2.group(1)))
                if linked and hasattr(linked, "as_ingredient_cached"):
                    linked = None
                ing = Ingredient(name=name, code=clean(rv[2]) if rv[2] is not None else "", unit=unit,
                                 pack_size=D(rv[7]) or Decimal("1"), vat_rate=D(rv[10]) or Decimal("0"))
                if isinstance(ing.code, str) and ing.code.endswith(".0"):
                    ing.code = ing.code[:-2]
                if linked:
                    ing.kind = Ingredient.SUBRECIPE
                    ing.recipe = linked
                    linked.as_ingredient_cached = True
                    ing.pack_size = Decimal("1")
                    ing.code = ""
                else:
                    ing.kind = Ingredient.PACKAGING if group.lower() == "packaging" else Ingredient.BOUGHT
                    ing.supplier = supplier(group)
                    ing.pack_price = D(rv[3]) if isinstance(rv[3], (int, float)) else None
                    if group.lower() == "mother recipe":
                        report["warnings"].append(
                            f"'{name}' is marked Mother Recipe but has a typed price ({rv[3]}) – kept as a bought item")
                ing.save()
                names[name.lower()] = ing

            # ---- recipe lines
            n_lines = 0
            for idx, (rv, rf) in enumerate(zip(rd_v, rd_f)):
                rec = recipes_by_row.get(idx + 6)
                if not rec:
                    continue
                pos = 0
                for slot in range(25):
                    c = 11 + 3 * slot  # 0-based index of the name column (L = 11)
                    iname, qv, qf = clean(rv[c]), rv[c + 1], rf[c + 1]
                    if not iname:
                        continue
                    ing = names.get(iname.lower())
                    if not ing:
                        report["warnings"].append(f"{rec.name}: ingredient '{iname}' is not on the ingredient list – skipped")
                        continue
                    section = RecipeLine.DOUGH if slot < 15 else (RecipeLine.TOPPING if slot < 20 else RecipeLine.PACKAGING)
                    mode, qty = RecipeLine.FIXED, D(qv) if isinstance(qv, (int, float)) else None
                    if isinstance(qf, str) and qf.startswith("="):
                        if re.search(r"\$D\$?\d+", qf):
                            mf = re.search(r"\$D\$?\d+\s*\*\s*([\d.]+)", qf)
                            mode, qty = RecipeLine.PER_UNIT, D(mf.group(1)) if mf else Decimal("1")
                        elif re.search(r"\$G\$?\d+", qf):
                            mf = re.search(r"\$G\$?\d+\s*\*\s*([\d.]+)", qf)
                            if mf:
                                mode, qty = RecipeLine.PER_KG, D(mf.group(1))
                    if qty is None:
                        report["warnings"].append(f"{rec.name}: '{iname}' has no quantity – skipped")
                        continue
                    pos += 1
                    RecipeLine.objects.create(recipe=rec, section=section, ingredient=ing, qty=qty, qty_mode=mode,
                                              position=pos)
                    n_lines += 1

            # ---- check against Excel's own batch cost (column CI)
            calc = Calculator().preload()
            diffs = []
            for idx, rv in enumerate(rd_v):
                rec = recipes_by_row.get(idx + 6)
                if not rec or not isinstance(rv[86], (int, float)):
                    continue
                ours = calc.batch_cost(rec)
                if abs(float(ours) - rv[86]) > 0.005:
                    diffs.append((rec.name, rv[86], float(ours)))

        self.stdout.write(self.style.SUCCESS(
            f"Imported {len(cats)} categories, {len(sups)} suppliers, {len(names)} ingredients, "
            f"{len(recipes_by_id)} recipes, {n_lines} recipe lines."))
        for w in report["warnings"]:
            self.stdout.write(self.style.WARNING("  ! " + w))
        if diffs:
            self.stdout.write(self.style.WARNING(f"{len(diffs)} recipe(s) cost differently from Excel:"))
            for n, a, b in diffs[:40]:
                self.stdout.write(f"   {n}: Excel {a:.4f}  program {b:.4f}")
        else:
            self.stdout.write(self.style.SUCCESS("Every recipe's batch cost matches the Excel file."))
